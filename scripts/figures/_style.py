"""
_style.py — shared python-pptx drawing helpers for the figure builders (BioSystems standard).

Page 6.7" x 9.5" portrait, Arial; panel label 12 pt bold, panel title 11 pt bold, axis titles 10 pt,
labels 8 pt (7 pt minimum); data lines 1.5 pt, axes 0.75 pt; Okabe-Ito colour-blind-safe palette.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "figures" / "source"
W, H, M = 6.7, 9.5, 0.35
DARK, GREY, LIGHT, WHITE = (RGBColor(0x22, 0x22, 0x22), RGBColor(0x55, 0x55, 0x55),
                            RGBColor(0xDD, 0xDD, 0xDD), RGBColor(0xFF, 0xFF, 0xFF))
OI = {"orange": "E69F00", "sky": "56B4E9", "green": "009E73", "yellow": "F0E442", "blue": "0072B2",
      "vermillion": "D55E00", "purple": "CC79A7", "grey": "999999", "black": "000000"}
SYSTEM_COLOURS = {
    "Gastrointestinal": "0072B2", "Genitourinary": "009E73", "Gynecologic/Breast": "CC79A7",
    "Hematologic": "D55E00", "Thoracic": "56B4E9", "Skin": "E69F00", "Endocrine": "999933", "CNS": "555555",
}
CAT_COLOURS = {"A": "0072B2", "B": "E69F00", "C": "009E73", "D": "CC79A7"}


def rgb(c) -> RGBColor:
    return c if isinstance(c, RGBColor) else RGBColor.from_string(c)


def new_slide():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


def text(slide, x, y, w, h, s, pt=8, bold=False, colour=DARK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE,
         italic=False, wrap=False, rotate=0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, line in enumerate(str(s).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        f = r.font
        f.name, f.size, f.bold, f.italic, f.color.rgb = "Arial", Pt(pt), bold, italic, rgb(colour)
    if rotate:
        tb.rotation = rotate
    return tb


def rich(slide, x, y, w, h, runs, pt=8, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, wrap=True):
    """runs: list of (text, dict(bold=, italic=, colour=)); '\n' inside a run text starts a paragraph."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    for t, fmt in runs:
        parts = t.split("\n")
        for k, part in enumerate(parts):
            if k:
                p = tf.add_paragraph()
                p.alignment = align
            if part:
                r = p.add_run()
                r.text = part
                f = r.font
                f.name, f.size = "Arial", Pt(fmt.get("pt", pt))
                f.bold, f.italic = fmt.get("bold", False), fmt.get("italic", False)
                f.color.rgb = rgb(fmt.get("colour", DARK))
    return tb


def rect(slide, x, y, w, h, fill, line=None, lw=0.5, shape=MSO_SHAPE.RECTANGLE):
    r = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        r.fill.background()
    else:
        r.fill.solid()
        r.fill.fore_color.rgb = rgb(fill)
    if line is None:
        r.line.fill.background()
    else:
        r.line.color.rgb = rgb(line)
        r.line.width = Pt(lw)
    r.shadow.inherit = False
    return r


def line(slide, x1, y1, x2, y2, colour=DARK, lw=0.75, dash=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(colour)
    c.line.width = Pt(lw)
    if dash:
        c.line.dash_style = dash
    return c


def panel(slide, x, y, letter, title, w=6.0):
    rich(slide, x, y, w, 0.25, [(f"({letter}) ", {"pt": 12, "bold": True}), (title, {"pt": 11, "bold": True})],
         wrap=False)


def lerp(lo, hi, t) -> RGBColor:
    t = max(0.0, min(1.0, t))
    lo, hi = rgb(lo), rgb(hi)
    return RGBColor(*[round(a + (b - a) * t) for a, b in zip(lo, hi)])


def blues(t: float) -> RGBColor:
    return lerp("F4F7FB", "08306B", t)


def ink(bg: RGBColor) -> RGBColor:
    r, g, b = bg
    return WHITE if 0.299 * r + 0.587 * g + 0.114 * b < 140 else DARK


def axes(slide, x, y, w, h, xlim, ylim, xticks, yticks, xlabel="", ylabel="", fmt_x=str, fmt_y=str, grid=True):
    """Draw a plain axis frame; returns a mapper (vx, vy) -> (inches x, inches y)."""
    def X(v):
        return x + (v - xlim[0]) / (xlim[1] - xlim[0]) * w

    def Y(v):
        return y + h - (v - ylim[0]) / (ylim[1] - ylim[0]) * h

    for t in yticks:
        if grid:
            line(slide, x, Y(t), x + w, Y(t), LIGHT, 0.5)
        text(slide, x - 0.55, Y(t) - 0.08, 0.5, 0.16, fmt_y(t), pt=8, align=PP_ALIGN.RIGHT)
    for t in xticks:
        text(slide, X(t) - 0.3, y + h + 0.04, 0.6, 0.16, fmt_x(t), pt=8, align=PP_ALIGN.CENTER)
    line(slide, x, y, x, y + h, DARK, 0.75)
    line(slide, x, y + h, x + w, y + h, DARK, 0.75)
    if xlabel:
        text(slide, x, y + h + 0.24, w, 0.18, xlabel, pt=10, align=PP_ALIGN.CENTER)
    if ylabel:
        text(slide, x - 0.55 - 0.1 - h / 2, y + h / 2 - 0.09, h, 0.18, ylabel, pt=10, align=PP_ALIGN.CENTER, rotate=270)
    return lambda vx, vy: (X(vx), Y(vy))


def marker(slide, cx, cy, size, fill, shape=MSO_SHAPE.OVAL, line_colour=None):
    return rect(slide, cx - size / 2, cy - size / 2, size, size, fill, line=line_colour, shape=shape)


def save(prs, name: str) -> Path:
    out = SRC / name
    prs.save(out)
    print(f"Saved {out.relative_to(ROOT)}")
    return out


def polygon(slide, pts, fill, alpha: float | None = None):
    """Filled polygon from a list of (x, y) inches; optional fill transparency (0-1 opacity)."""
    fb = slide.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]), scale=1.0)
    fb.add_line_segments([(Inches(x), Inches(y)) for x, y in pts[1:]], close=True)
    shp = fb.convert_to_shape()
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    shp.line.fill.background()
    shp.shadow.inherit = False
    if alpha is not None:
        from pptx.oxml.ns import qn
        clr = shp.fill._xPr.find(qn("a:solidFill"))[0]
        a = clr.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))})
        clr.append(a)
    return shp


def polyline(slide, pts, colour, lw=1.5):
    for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
        line(slide, x1, y1, x2, y2, colour, lw)
